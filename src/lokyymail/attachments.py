"""Text aus Anhängen ziehen: PDF, DOCX, HTML, Text/CSV/JSON. Nichts wird ausgeführt, nur gelesen.

Herkunft: übernommen aus Gmail Guard (gleicher Urheber), ergänzt um Schutz gegen Zip-Bomben bei DOCX
und um unsere Behandlung von verstecktem Text in HTML-Anhängen.
"""

from __future__ import annotations

import io
import zipfile

from .sanitize import ai_text

TEXT_EXT = (".txt", ".csv", ".md", ".json", ".xml", ".ics", ".log", ".eml", ".vcf")
HTML_EXT = (".html", ".htm")
MAX_PDF_PAGES = 100
MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024
MAX_CHARS = 40_000


def extract_text(data: bytes, filename: str, mime: str) -> tuple[str | None, str]:
    """Gibt (text, hinweis) zurück. text=None, wenn das Format nicht unterstützt oder nicht lesbar ist."""
    name = (filename or "").lower()
    mime = (mime or "").lower()

    if mime == "application/pdf" or name.endswith(".pdf"):
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError

        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception:
                    return None, "PDF ist passwortgeschützt."
            total = len(reader.pages)
            parts, found_text = [], False
            for i, page in enumerate(reader.pages[:MAX_PDF_PAGES], 1):
                try:
                    page_text = page.extract_text() or ""
                except Exception:
                    page_text = "[nicht lesbar]"
                found_text = found_text or bool(page_text.strip())
                parts.append(f"--- Seite {i} ---\n{page_text}")
        except (PdfReadError, ValueError, KeyError) as exc:
            return None, f"PDF ist beschädigt ({type(exc).__name__})."
        hint = f"Nur die ersten {MAX_PDF_PAGES} von {total} Seiten gelesen." if total > MAX_PDF_PAGES else ""
        if not found_text:
            hint = (hint + " Kein Text gefunden, vermutlich ein Scan (Bild-PDF).").strip()
        return "\n".join(parts), hint

    if name.endswith(".docx") or mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                if sum(i.file_size for i in zf.infolist()) > MAX_DOCX_UNCOMPRESSED:
                    return None, "DOCX ist entpackt zu groß (Schutz vor Zip-Bomben)."
            from docx import Document

            doc = Document(io.BytesIO(data))
        except Exception as exc:
            return None, f"DOCX ist beschädigt ({type(exc).__name__})."
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts), ""

    if name.endswith(HTML_EXT) or mime == "text/html":
        result = ai_text(html=data.decode("utf-8", errors="replace"), plain=None, limit=MAX_CHARS)
        hint = "Versteckter Text im HTML wurde entfernt." if result.hidden_text_chars else ""
        return result.text, hint

    if name.endswith(TEXT_EXT) or mime.startswith("text/"):
        return data.decode("utf-8", errors="replace"), ""

    return None, f"Format nicht unterstützt ({mime or 'unbekannt'}). Unterstützt: PDF, DOCX, HTML, Text/CSV/JSON."
